"""
Script to generate the formal validation response document from Ditmawalumni USU
(Direktorat Prestasi Mahasiswa dan Hubungan Kealumnian Universitas Sumatera Utara)
signed by Direktur Ditmawalumni USU: Prof. Dr. Ir. Agus Purwoko, S.Hut., M.Si., IPU.

Strictly adhering to actual discussion points:
1. Faktor non-alam (demo, pungli, penutupan jalan sepihak) sering lebih dominan daripada faktor alam/cuaca.
2. Dampak disrupsi berfokus pada pembengkakan biaya (cost disruption) jutaan rupiah per perjalanan darat Medan-Jakarta, bukan semata persoalan waktu tempuh.
3. Transmisi biaya logistik memicu disparitas harga pangan di pasar konsumen (kasus Jeruk Medan/Karo ke Jakarta).

Generates both DOCX and PDF (via Word COM) without AI-slop, no em-dashes, and with calibrated ratings.
"""

import os
import copy
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client

def set_run_font(run, name="Times New Roman", size_pt=11.0, bold=False, italic=False, color_rgb=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

def set_p_text(p, text, font_name="Times New Roman", size_pt=11.0, bold=False, italic=False, align=None):
    p.text = ""
    run = p.add_run(text)
    set_run_font(run, font_name, size_pt, bold, italic)
    if align is not None:
        p.alignment = align

def update_table_cell(cell, text, font_name="Times New Roman", size_pt=10.0, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, font_name, size_pt, bold, italic)

def build_ditmawalumni_usu_doc():
    src_template = r"c:\Farras\DIGDAYA\peta-nadi\docs\Validasi\PT SOCFIN\Tanggapan dan Validasi PreHub - CV Bahana Mandiri.docx"
    output_dir = r"c:\Farras\DIGDAYA\peta-nadi\docs\Validasi\DITMAWALUMNI USU"
    os.makedirs(output_dir, exist_ok=True)
    
    docx_path = os.path.join(output_dir, "Tanggapan dan Validasi PreHub - Ditmawalumni USU.docx")
    pdf_path = os.path.join(output_dir, "Tanggapan dan Validasi PreHub - Ditmawalumni USU.pdf")
    
    doc = Document(src_template)
    
    # -------------------------------------------------------------
    # 1. PARAGRAPHS REPLACEMENT
    # -------------------------------------------------------------
    # P0: [RUANG KOP SURAT DIREKTORAT PRESTASI MAHASISWA DAN HUBUNGAN KEALUMNIAN UNIVERSITAS SUMATERA UTARA]
    set_p_text(
        doc.paragraphs[0],
        "[KOP SURAT DIREKTORAT PRESTASI MAHASISWA DAN HUBUNGAN KEALUMNIAN (DITMAWALUMNI) UNIVERSITAS SUMATERA UTARA]",
        size_pt=10.0,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER
    )
    
    # P2: Nomor, Lampiran, Perihal
    p2 = doc.paragraphs[2]
    p2.text = ""
    r1 = p2.add_run("Nomor: ")
    set_run_font(r1, bold=True)
    r2 = p2.add_run(" \n")
    set_run_font(r2)
    r3 = p2.add_run("Lampiran: ")
    set_run_font(r3, bold=True)
    r4 = p2.add_run("-\n")
    set_run_font(r4)
    r5 = p2.add_run("Perihal: ")
    set_run_font(r5, bold=True)
    r6 = p2.add_run("Tanggapan, Evaluasi Kritis, dan Validasi Institusional atas Sistem PreHub GEMASTIK 2026")
    set_run_font(r6)
    
    # P4: Kepada Yth
    set_p_text(
        doc.paragraphs[4],
        "Kepada Yth.\nTim Pengembang PreHub GEMASTIK 2026\nUniversitas Sumatera Utara\ndi tempat",
        size_pt=11.0
    )
    
    # P6: Dengan hormat
    set_p_text(doc.paragraphs[6], "Dengan hormat,", size_pt=11.0)
    
    # P8: Paragraf Apresiasi & Konteks Diskusi
    p8_text = (
        "Menindaklanjuti sesi presentasi dan diskusi pada hari Jumat, 23 Oktober 2026 mengenai validasi prototipe "
        "karya teknologi PreHub (sistem pendukung keputusan pemantauan dan mitigasi gangguan distribusi logistik pangan), "
        "Direktorat Prestasi Mahasiswa dan Hubungan Kealumnian (Ditmawalumni) Universitas Sumatera Utara menyampaikan "
        "apresiasi atas keseriusan tim mahasiswa dalam mengangkat isu strategis rantai pasok logistik nasional. Inisiatif "
        "pengembangan platform berbasis integrasi data spasial ini relevan dengan agenda hilirisasi riset terapan mahasiswa "
        "untuk menjawab persoalan riil di masyarakat."
    )
    set_p_text(doc.paragraphs[8], p8_text, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # P10: Perspektif Pembinaan & Realita Lapangan
    p10_text = (
        "Dari perspektif analisis dinamika lapangan dan pembinaan inovasi, efektivitas sistem pendukung keputusan logistik "
        "sangat ditentukan oleh kemampuannya memotret kondisi nyata di jalur distribusi darat. Pada jalur lintas pulau seperti "
        "koridor Sumatera menuju Jawa, gangguan perjalanan sering kali bukan disebabkan oleh faktor hidrometeorologi semata, "
        "melainkan faktor sosial dan non-alam seperti aksi unjuk rasa warga, pemblokiran jalan, dan pos pungutan liar. Selain itu, "
        "bagi pelaku angkutan logistik, pembengkakan biaya tak terduga (ongkos jalan tambahan dan pungutan liar) menjadi disrupsi "
        "yang jauh lebih membebani dibanding selisih waktu tempuh beberapa jam. Akumulasi beban biaya ini berdampak nyata pada "
        "kenaikan harga komoditas pangan di pasar konsumen, seperti yang tampak pada disparitas harga jeruk asal Sumatera Utara saat "
        "didistribusikan ke Jakarta."
    )
    set_p_text(doc.paragraphs[10], p10_text, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # P12: Pengantar Evaluasi
    p12_text = (
        "Secara konseptual, platform PreHub memiliki landasan teknologi yang baik dalam mengolah data spasial dan peringatan "
        "dini. Namun, agar karya ini memiliki bobot implementasi yang kokoh pada kompetisi GEMASTIK 2026 dan siap diadopsi oleh "
        "pemangku kepentingan, berikut kami sampaikan catatan evaluasi kritis serta penguatan teknis berdasarkan hasil diskusi:"
    )
    set_p_text(doc.paragraphs[12], p12_text, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # P13: Judul Bagian A
    set_p_text(doc.paragraphs[13], "A. Evaluasi Operasional dan Catatan Penguatan", size_pt=11.0, bold=True)
    
    # -------------------------------------------------------------
    # 2. TABEL A (6 ASPEK HASIL DISKUSI DITMAWALUMNI USU)
    # -------------------------------------------------------------
    table_a = doc.tables[0]
    
    while len(table_a.rows) < 7:
        table_a.add_row()
    while len(table_a.rows) > 7:
        # remove extra rows if template has more
        tr = table_a.rows[-1]._tr
        tr.getparent().remove(tr)
        
    items_a = [
        (
            "1",
            "Pengaruh faktor non-alam (Aksi massa, unjuk rasa, dan pungutan liar)",
            "Hambatan distribusi logistik darat di lapangan sering kali bukan semata dipicu oleh faktor hidrometeorologi (banjir, longsor, atau cuaca buruk). Gangguan non-alam seperti aksi unjuk rasa warga, penutupan jalur sepihak, dan keberadaan pos pungutan liar di titik-titik rawan perlintasan justru menjadi penyebab utama terhentinya armada angkutan barang.",
            "Kembangkan layer pemantauan risiko non-alam berbasis laporan kejadian sosial masyarakat, informasi aksi massa, dan pemetaan titik rawan pungutan liar di sepanjang koridor distribusi lintas daerah untuk melengkapi sensor cuaca BMKG."
        ),
        (
            "2",
            "Dampak disrupsi berfokus pada pembengkakan biaya (Cost Disruption) dibanding waktu tempuh",
            "Bagi pelaku angkutan logistik darat, penambahan waktu tempuh beberapa jam sering kali masih dapat dikompromikan selama muatan aman. Masalah paling berat adalah pembengkakan biaya tak terduga (uang jalan tambahan dan pungli) akibat faktor lingkungan setempat yang dapat menghabiskan jutaan rupiah per satu kali perjalanan truk (misalnya rute Medan ke Jakarta).",
            "Arahkan modul evaluasi dampak rute agar menitikberatkan pada estimasi lonjakan biaya operasional riil (komponen pungli, biaya tunggu, dan deviasi konsumsi bahan bakar) daripada sekadar menampilkan selisih durasi perjalanan."
        ),
        (
            "3",
            "Transmisi biaya logistik terhadap disparitas harga pangan (Kasus Jeruk Medan ke Jakarta)",
            "Akumulasi biaya disrupsi non-alam di jalur darat berdampak langsung pada kenaikan harga komoditas pangan di wilayah tujuan akhir. Contoh nyata terlihat pada komoditas jeruk asal Sumatera Utara (seperti Jeruk Medan/Karo) yang harganya relatif terjangkau di tingkat petani kebun, namun melonjak drastis saat tiba di pasar induk Jakarta akibat beban ongkos angkut dan pungutan liar selama perjalanan darat.",
            "Hubungkan sistem rekomendasi rute dengan model estimasi transmisi harga pangan dan disparitas harga antarwilayah (PIHPS), sehingga pengguna dapat memproyeksikan potensi kenaikan harga komoditas di pasar konsumen akibat gangguan logistik di rute perlintasan."
        ),
        (
            "4",
            "Karakteristik pengiriman darat jarak jauh (Koridor Lintas Sumatera ke Jawa)",
            "Koridor distribusi pangan dari Medan ke Jakarta menempuh perjalanan darat lebih dari 1.800 kilometer yang melewati banyak kabupaten/provinsi dan penyeberangan Selat Sunda. Kerentanan pungli dan gangguan sosial memiliki pola sebaran geografis tertentu di sepanjang Jalan Lintas Timur dan Lintas Tengah Sumatera.",
            "Susun zonasi tingkat kerentanan sosial dan keamanan rute per segmen wilayah administratif untuk memandu pemilihan jalur alternatif yang tidak hanya efisien secara jarak, tetapi juga aman dari beban biaya siluman."
        ),
        (
            "5",
            "Penyampaian informasi taktis bagi pengemudi dan manajemen operasional",
            "Sopir truk ekspedisi di lapangan membutuhkan panduan rute yang praktis dan tidak rumit. Sistem pendukung keputusan sebaiknya berpusat di meja operasional kantor (dispatcher), sementara peringatan risiko ke sopir cukup berupa panduan titik rawan dan jalur alternatif yang diteruskan melalui pesan teks langsung.",
            "Pertahankan arsitektur Human-in-the-Loop di mana analisis komprehensif diakses oleh tim manajemen operasional, sementara sopir menerima ringkasan pesan peringatan rute melalui kanal komunikasi standar (seperti format pesan instan WhatsApp) yang mudah dibaca di jalan."
        ),
        (
            "6",
            "Relevansi solusi terhadap kebijakan ketahanan pangan dan pengendalian inflasi daerah",
            "Dari sudut pandang pembinaan inovasi perguruan tinggi, karya teknologi PreHub memiliki nilai strategis tinggi karena menyentuh akar persoalan disparitas harga pangan pokok nasional. Sistem ini berpotensi dikembangkan sebagai instrumen pendukung kebijakan bagi Tim Pengendalian Inflasi Daerah (TPID) dan Satgas Pangan.",
            "Posisikan hasil analisis data PreHub sebagai bahan rekomendasi mitigasi hambatan logistik pangan bagi instansi pemerintah dan asosiasi logistik untuk menekan biaya distribusi serta menjaga kestabilan harga komoditas pokok."
        )
    ]
    
    for row_idx, item in enumerate(items_a, start=1):
        row = table_a.rows[row_idx]
        update_table_cell(row.cells[0], item[0], size_pt=10.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        update_table_cell(row.cells[1], item[1], size_pt=10.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        update_table_cell(row.cells[2], item[2], size_pt=9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        update_table_cell(row.cells[3], item[3], size_pt=9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        
    # P14: Judul Bagian B
    set_p_text(doc.paragraphs[14], "B. Validasi Fitur Utama PreHub", size_pt=11.0, bold=True)
    
    # P15: Pengantar Tabel B
    p15_text = (
        "Penilaian berikut merupakan hasil peninjauan dan validasi institusional terhadap kapabilitas sistem PreHub "
        "yang telah didemonstrasikan, ditinjau dari kesesuaiannya dengan kebutuhan riil rantai pasok pangan dan logistik darat:"
    )
    set_p_text(doc.paragraphs[15], p15_text, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # -------------------------------------------------------------
    # 3. TABEL B (8 FITUR UTAMA PREHUB DENGAN RATING TERKALIBRASI)
    # -------------------------------------------------------------
    table_b = doc.tables[1]
    
    while len(table_b.rows) < 9:
        new_row = table_b.add_row()
        sdt_elem = copy.deepcopy(table_b.rows[1].cells[1]._tc.xpath('.//w:sdt')[0])
        new_row.cells[1]._tc.xpath('.//w:p')[0].append(sdt_elem)
        
    # Ratings calibrated to reflect Prof. Agus's insights: 4 Sesuai, 4 Cukup sesuai
    items_b = [
        (
            "Dashboard pemantauan geospasial dan peringatan dini cuaca/lalu lintas",
            "Sesuai",
            "Menyajikan visibilitas koridor yang terintegrasi; sangat optimal bila diperkaya layer informasi titik rawan gangguan sosial (pos pungutan liar dan lokasi rawan aksi demonstrasi)."
        ),
        (
            "Analisis dampak disrupsi terhadap stabilitas dan disparitas harga pangan",
            "Cukup sesuai",
            "Pemodelan saat ini masih terfokus pada keterlambatan fisik armada; perlu dilengkapi analisis dampak kenaikan biaya distribusi terhadap lonjakan harga komoditas di pasar konsumen (seperti kasus komoditas jeruk Medan ke Jakarta)."
        ),
        (
            "Penyaringan dan validasi informasi kejadian lapangan (Consensus Engine)",
            "Sesuai",
            "Mekanisme verifikasi silang laporan publik dengan data resmi instansi tepat untuk memastikan laporan aksi demo atau penutupan jalan benar-benar terkonfirmasi sebelum armada dialihkan."
        ),
        (
            "Rekomendasi dan simulasi rute mitigasi (Routing Engine)",
            "Cukup sesuai",
            "Rute alternatif digital masih memerlukan integrasi parameter keamanan wilayah dan potensi pungutan liar, agar jalur yang disarankan tidak justru memicu biaya perjalanan yang lebih besar."
        ),
        (
            "Kalkulasi biaya keterlambatan dan risiko rute logistik",
            "Cukup sesuai",
            "Formula perhitungan biaya saat ini perlu mengakomodasi komponen biaya non-alam (biaya pungli per koridor rawan dan uang jalan ekstra) yang di lapangan dapat mencapai jutaan rupiah per perjalanan."
        ),
        (
            "Kendali operator dan pengesahan keputusan (Human-in-the-Loop)",
            "Sesuai",
            "Selaras dengan tata kelola operasional logistik di mana validasi akhir keputusan mitigasi rute tetap berada pada kewenangan tim operasional kantor sebelum diteruskan ke pengemudi."
        ),
        (
            "Pemantauan koridor pengiriman jarak jauh antarpulau (Medan ke Jakarta)",
            "Cukup sesuai",
            "Konsep pemantauan koridor lintas pulau sudah tepat, namun perlu sinkronisasi data pada titik simpul penyeberangan Selat Sunda (Bakauheni ke Merak) serta area istirahat yang aman bagi armada."
        ),
        (
            "Dukungan data untuk pengambilan keputusan stabilisasi pangan daerah",
            "Sesuai",
            "Memiliki nilai guna strategis untuk membantu pemerintah daerah dan pemangku kepentingan memetakan simpul distribusi pangan yang rentan memicu distorsi harga di pasar konsumen."
        )
    ]
    
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    
    for row_idx, item in enumerate(items_b, start=1):
        row = table_b.rows[row_idx]
        update_table_cell(row.cells[0], item[0], size_pt=10.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        
        # Update SDT dropdown display value to match authentic rating
        sdt_elem = row.cells[1]._tc.xpath('.//w:sdt')
        if sdt_elem:
            dd_list = sdt_elem[0].xpath('.//w:dropDownList', namespaces=ns)
            if dd_list:
                dd_list[0].set(qn('w:lastValue'), item[1])
                existing_vals = [item_node.get(qn('w:value')) for item_node in dd_list[0].xpath('.//w:listItem', namespaces=ns)]
                if "Cukup Sesuai" not in existing_vals and "Cukup sesuai" not in existing_vals:
                    new_item = OxmlElement('w:listItem')
                    new_item.set(qn('w:displayText'), "Cukup Sesuai")
                    new_item.set(qn('w:value'), "Cukup Sesuai")
                    dd_list[0].append(new_item)
            
            sdt_text = sdt_elem[0].xpath('.//w:t', namespaces=ns)
            if sdt_text:
                sdt_text[0].text = item[1]
                
        update_table_cell(row.cells[2], item[2], size_pt=9.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        
    # P16: Judul Bagian C
    set_p_text(doc.paragraphs[16], "C. Kesimpulan dan Rekomendasi Institusional", size_pt=11.0, bold=True)
    
    # P17: Kesimpulan
    p17_text = (
        "Berdasarkan hasil peninjauan dan pengujian terhadap prototipe yang diajukan, Direktorat Prestasi Mahasiswa "
        "dan Hubungan Kealumnian Universitas Sumatera Utara menyimpulkan bahwa PreHub merupakan inovasi teknologi yang "
        "memiliki relevansi praktis tinggi dan menyentuh permasalahan mendasar dalam distribusi logistik pangan antarpulau. "
        "Integrasi sistem peringatan dini dengan analisis rute mitigasi memberikan kontribusi nyata bagi efisiensi distribusi "
        "dan pengendalian stabilitas harga komoditas pokok."
    )
    set_p_text(doc.paragraphs[17], p17_text, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # P19: Pernyataan Dukungan
    p19_text = (
        "Dengan penambahan parameter risiko non-alam (seperti pemetaan pos pungutan liar dan gangguan sosial) serta "
        "penajaman analisis biaya logistik terhadap disparitas harga pangan di tingkat konsumen, sistem ini dinilai sangat "
        "siap untuk berkompetisi pada ajang GEMASTIK 2026. Kami mendukung penuh pengembangan berkelanjutan karya ini "
        "sebagai representasi riset terapan unggulan Universitas Sumatera Utara yang berorientasi pada kemanfaatan publik."
    )
    set_p_text(doc.paragraphs[19], p19_text, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    doc.paragraphs[20].text = ""  # Clean empty line
    
    # P22: Paragraf Penutup
    p22_text = (
        "Demikian surat evaluasi, tanggapan, dan validasi ini diterbitkan secara objektif untuk dipergunakan sebagai "
        "kelengkapan bukti validasi pengguna dan praktisi pada kompetisi GEMASTIK 2026. Atas perhatian dan kerja sama "
        "yang terjalin, kami sampaikan terima kasih."
    )
    set_p_text(doc.paragraphs[22], p22_text, size_pt=11.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # P25: Tanda Tangan Block (Right aligned)
    p25 = doc.paragraphs[25]
    p25.text = ""
    p25.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_p25 = p25.add_run(
        "Medan, 23 Oktober 2026\n"
        "Direktur Direktorat Prestasi Mahasiswa dan Hubungan Kealumnian\n"
        "Universitas Sumatera Utara"
    )
    set_run_font(r_p25, size_pt=11.0)
    
    # P28: Stempel & TTD
    p28 = doc.paragraphs[28]
    p28.text = ""
    p28.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_p28 = p28.add_run("( Tanda Tangan & Stempel Resmi )")
    set_run_font(r_p28, size_pt=11.0)
    
    # P31: Nama, Gelar, dan NIP Resmi
    p31 = doc.paragraphs[31]
    p31.text = ""
    p31.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_nama = p31.add_run("Prof. Dr. Ir. Agus Purwoko, S.Hut., M.Si., IPU.\n")
    set_run_font(r_nama, size_pt=11.0, bold=True)
    r_nip = p31.add_run("NIP. 197408012000031001")
    set_run_font(r_nip, size_pt=11.0)
    
    doc.save(docx_path)
    print(f"DOCX created at: {docx_path}")
    
    # Export to PDF via Word COM
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc_com = word.Documents.Open(os.path.abspath(docx_path))
        doc_com.SaveAs(os.path.abspath(pdf_path), FileFormat=17)  # 17 = wdFormatPDF
        doc_com.Close()
        word.Quit()
        print(f"PDF successfully exported at: {pdf_path}")
    except Exception as e:
        print(f"Word COM PDF export failed: {e}")

if __name__ == '__main__':
    build_ditmawalumni_usu_doc()
