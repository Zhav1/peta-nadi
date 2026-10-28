# -*- coding: utf-8 -*-
"""Canonical builder for the PreHub internal validation sheet.

Renders `Final-Gemastik/Lembar_Validasi_Internal_PreHub.md` into:

  - Final-Gemastik/Lembar_Validasi_Internal_PreHub.docx
  - Final-Gemastik/Lampiran_N60_PreHub.docx   (byte-identical copy)
  - Final-Gemastik/Lembar_Validasi_Internal_PreHub.pdf  (Word COM, LibreOffice fallback)
  - Final-Gemastik/Lampiran_N60_PreHub.pdf    (byte-identical copy)
  - Final-Gemastik/validation_doc_manifest.json

The markdown file is the single source of truth: every heading, paragraph,
table row, and wording in the DOCX/PDF comes from it, so the two formats cannot
drift. Only presentation is decided here (page geometry, type scale, colours,
header/footer, and the landscape section that carries the 60-row table).

No network calls. The build stops if the markdown is missing or if a table row
does not match its header width.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "Final-Gemastik"
MD_PATH = OUT_DIR / "Lembar_Validasi_Internal_PreHub.md"

DOCX_CANON = OUT_DIR / "Lembar_Validasi_Internal_PreHub.docx"
DOCX_LAMP = OUT_DIR / "Lampiran_N60_PreHub.docx"
PDF_CANON = OUT_DIR / "Lembar_Validasi_Internal_PreHub.pdf"
PDF_LAMP = OUT_DIR / "Lampiran_N60_PreHub.pdf"
MANIFEST_PATH = OUT_DIR / "validation_doc_manifest.json"

DOC_TITLE = "PreHub: Lembar Validasi Internal"
DOC_DATE = "28 Oktober 2026"

PORTRAIT_WIDTH = 17.2
LANDSCAPE_WIDTH = 26.1

NAVY = "1F3864"
TEAL = "2E7D6E"
LIGHT = "EAF1F8"
GRAY = "F2F2F2"
GREEN = "E2EFDA"
AMBER = "FFF2CC"
RED = "FBE2E2"
WHITE = "FFFFFF"
MID = "595959"
BORDER = "BFBFBF"
FONT = "Calibri"

SEPARATOR = re.compile(r"^:?-{3,}:?$")


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


# ---------------------------------------------------------------------------
# DOCX primitives
# ---------------------------------------------------------------------------


def shade(cell, color: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    for old in tc.findall(qn("w:shd")):
        tc.remove(old)
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:color"), "auto")
    element.set(qn("w:fill"), color)
    tc.append(element)


def borders(table) -> None:
    properties = table._tbl.tblPr
    for old in properties.findall(qn("w:tblBorders")):
        properties.remove(old)
    block = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)
        block.append(element)
    properties.append(block)


def no_table_borders(table) -> None:
    properties = table._tbl.tblPr
    for old in properties.findall(qn("w:tblBorders")):
        properties.remove(old)
    block = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        block.append(element)
    properties.append(block)


def cell_margins(table, top: int = 30, left: int = 70, bottom: int = 30, right: int = 70) -> None:
    properties = table._tbl.tblPr
    for old in properties.findall(qn("w:tblCellMar")):
        properties.remove(old)
    block = OxmlElement("w:tblCellMar")
    for side, width in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")
        block.append(element)
    properties.append(block)


def fixed(table, widths: Sequence[float]) -> None:
    properties = table._tbl.tblPr
    for old in properties.findall(qn("w:tblLayout")):
        properties.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    for row in table.rows:
        for index, width in enumerate(widths):
            if index < len(row.cells):
                row.cells[index].width = Cm(width)


def no_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(element)


def add_field(paragraph, code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def add_rich(paragraph, text: str, size: float, bold: bool = False, color: Optional[str] = None) -> None:
    """Add text, honouring `**bold**` spans carried over from the markdown."""
    for index, chunk in enumerate(str(text).split("**")):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        run.font.size = Pt(size)
        run.bold = bold or (bool(index % 2))
        if color:
            run.font.color.rgb = rgb(color)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths: Sequence[float],
    size: float,
    header_size: Optional[float] = None,
    center: Optional[set] = None,
    first_bold: bool = False,
    zebra: bool = True,
    row_fills: Optional[Dict[int, str]] = None,
    margin: Optional[Tuple[int, int, int, int]] = None,
    keep_rows_together: bool = False,
) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    if margin:
        cell_margins(table, *margin)
    else:
        cell_margins(table)

    header = table.rows[0]
    no_split(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        cell.text = ""
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_with_next = True
        add_rich(paragraph, label, header_size or size, bold=True, color=WHITE)
    repeat_header(header)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        no_split(row)
        row_fill = (row_fills or {}).get(row_index)
        for column_index, header_label in enumerate(headers):
            value = values[column_index] if column_index < len(values) else ""
            cell = row.cells[column_index]
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if center and column_index in center:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_rich(paragraph, value, size, bold=first_bold and column_index == 0)
            if row_fill:
                shade(cell, row_fill)
            elif zebra and row_index % 2 == 1:
                shade(cell, GRAY)

    fixed(table, widths)
    if keep_rows_together:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    return table


def para(
    doc: Document,
    text: str,
    size: float = 9.0,
    italic: bool = False,
    color: Optional[str] = None,
    space_after: float = 5,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    bold: bool = False,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.keep_together = True
    add_rich(paragraph, text, size, bold=bold, color=color)
    if italic:
        for run in paragraph.runs:
            run.italic = True


def section_band(doc: Document, number: str, title: str, width: float, size: float = 11.5) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    cell_margins(table, 80, 130, 80, 130)
    cell = table.cell(0, 0)
    shade(cell, NAVY)
    no_split(table.rows[0])
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = True
    number_run = paragraph.add_run(f"{number}  ")
    number_run.bold = True
    number_run.font.size = Pt(size - 2.3)
    number_run.font.color.rgb = rgb("9DC3E6")
    title_run = paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(size)
    title_run.font.color.rgb = rgb(WHITE)
    fixed(table, [width])


def spacer(doc: Document, points: float = 4.5, keep_next: bool = False) -> None:
    paragraph = doc.add_paragraph("")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(points)
    paragraph.paragraph_format.keep_with_next = keep_next


def subsection(doc: Document, text: str, size: float = 10.5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(9)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(TEAL)


def setup_header_footer(section, width: float, label: str) -> None:
    header = section.header
    header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    for paragraph in list(header.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    table = header.add_table(rows=1, cols=2, width=Cm(width))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell_margins(table, 0, 0, 0, 0)
    no_table_borders(table)
    fixed(table, [width * 0.66, width * 0.34])

    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    left.paragraph_format.space_before = Pt(0)
    run = left.add_run("PreHub · Lembar Validasi Internal")
    run.font.size = Pt(7.6)
    run.font.color.rgb = rgb(MID)

    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    right.paragraph_format.space_before = Pt(0)
    run = right.add_run(label)
    run.font.size = Pt(7.6)
    run.font.color.rgb = rgb(MID)

    tail = header.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Halaman ")
    run.font.size = Pt(7.6)
    run.font.color.rgb = rgb(MID)
    add_field(footer, "PAGE")
    run = footer.add_run(" dari ")
    run.font.size = Pt(7.6)
    run.font.color.rgb = rgb(MID)
    add_field(footer, "NUMPAGES")


def page_setup(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.3)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.4)
        section.left_margin = Cm(1.9)
        section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.65)


# ---------------------------------------------------------------------------
# Markdown parsing
# ---------------------------------------------------------------------------


def split_row(line: str) -> List[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_markdown(path: Path) -> List[Tuple[str, Any]]:
    """Turn the validation markdown into ('h1'|'h2'|'h3'|'p'|'table', payload) blocks."""
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: List[Tuple[str, Any]] = []
    index = 0
    total = len(lines)

    while index < total:
        text = lines[index].strip()
        if not text:
            index += 1
            continue

        if text.startswith("|"):
            raw_rows: List[List[str]] = []
            while index < total and lines[index].strip().startswith("|"):
                raw_rows.append(split_row(lines[index]))
                index += 1
            if len(raw_rows) < 2 or not all(SEPARATOR.match(cell or "") for cell in raw_rows[1]):
                raise SystemExit(f"Tabel markdown tanpa baris pemisah pada '{raw_rows[0][0] if raw_rows else '?'}'.")
            header = raw_rows[0]
            body = raw_rows[2:]
            for row in body:
                if len(row) != len(header):
                    raise SystemExit(
                        f"Baris tabel '{header[0]}' berisi {len(row)} sel, seharusnya {len(header)}."
                    )
            blocks.append(("table", (header, body)))
            continue

        if text.startswith("### "):
            blocks.append(("h3", text[4:].strip()))
            index += 1
            continue
        if text.startswith("## "):
            blocks.append(("h2", text[3:].strip()))
            index += 1
            continue
        if text.startswith("# "):
            blocks.append(("h1", text[2:].strip()))
            index += 1
            continue

        parts = [text]
        index += 1
        while index < total:
            following = lines[index].strip()
            if not following or following.startswith(("|", "#")):
                break
            parts.append(following)
            index += 1
        blocks.append(("p", " ".join(parts)))

    return blocks


# ---------------------------------------------------------------------------
# Table geometry
# ---------------------------------------------------------------------------

COLUMN_WIDTHS: Dict[str, List[float]] = {
    "Item|Keterangan": [3.7, 13.5],
    "Mitra|Tanggal|Cakupan instrumen|Hasil pada surat": [4.6, 3.0, 3.4, 6.2],
    "Tema fitur|PT MTI|CV Bahana|PT Socfin|Bumi Sawita|USU": [6.2, 2.1, 2.2, 2.1, 2.2, 2.4],
    "Kode|Pernyataan|Aspek PreHub yang diuji|Rata-rata jawaban|Poin|Kontribusi skor": [
        1.2,
        5.4,
        5.2,
        2.0,
        1.6,
        1.8,
    ],
    "Skenario|Fokus|Manual|PreHub|Pengurangan|Status": [1.6, 6.3, 2.1, 2.1, 2.2, 2.9],
    "Metrik|Hasil|Target|Status": [4.6, 6.6, 3.2, 2.8],
    "Metrik|Uji Skenario (N=60)|Kejadian Aktual (N=20)|Target|Status": [3.9, 4.5, 4.6, 2.2, 2.0],
    "Metrik|Nilai|Arti bagi operator": [3.4, 3.6, 10.2],
    "Metrik spasial (N=60)|Hasil": [11.2, 6.0],
    "Kanal|Bobot|Sinyal utama|Kondisi yang menekan skor": [3.5, 1.5, 6.6, 5.6],
    "ID|Kelas|Lokasi dan koridor|Skor konsensus vs ambang 50%|Aktual → Prediksi|Penyebab|Perbaikan": [
        1.9,
        0.8,
        3.8,
        2.4,
        2.0,
        3.3,
        3.0,
    ],
    "No|ID & Provinsi|Sumber & Tanggal|Skenario & Judul Berita|Koridor Aktual vs Lokasi Berita (Jarak)|"
    "Keyakinan Agen (Cuaca/Lalin/Berita/Harga)|Probabilitas Kalibrasi (%)|Aktual|Prediksi|Kelas|Brier Score": [
        0.6,
        2.2,
        2.5,
        7.9,
        4.2,
        2.6,
        1.55,
        1.4,
        1.4,
        0.85,
        1.0,
    ],
}

RESPONDENT_HEADERS = ("ID", "Profil", "Q1")
RESPONDENT_WIDTHS = [0.9, 4.9, 0.85, 0.85, 0.85, 0.85, 0.85, 0.85, 0.85, 0.85, 0.85, 0.85, 1.3, 1.6]

TABLE_SIZES = {
    11: 6.3,
    14: 6.4,
    7: 6.8,
    6: 7.5,
    5: 7.7,
    4: 7.7,
    3: 8.0,
    2: 8.3,
}


def column_widths(header: Sequence[str], body: Sequence[Sequence[str]], total: float) -> List[float]:
    key = "|".join(header)
    widths: Optional[List[float]] = None
    if header[:3] == list(RESPONDENT_HEADERS) and len(header) == 14:
        widths = list(RESPONDENT_WIDTHS)
    elif key in COLUMN_WIDTHS and len(COLUMN_WIDTHS[key]) == len(header):
        widths = list(COLUMN_WIDTHS[key])

    if widths is None:
        weights = []
        for index in range(len(header)):
            lengths = [len(row[index]) for row in body if index < len(row)]
            lengths.append(len(header[index]) * 2)
            weights.append(max(min(max(lengths), 90), 8))
        factor = total / float(sum(weights))
        widths = [weight * factor for weight in weights]

    scale = total / float(sum(widths))
    return [width * scale for width in widths]


def centered_columns(header: Sequence[str], body: Sequence[Sequence[str]]) -> set:
    center = set()
    for index in range(len(header)):
        values = [row[index].replace("**", "").strip() for row in body if index < len(row)]
        values = [value for value in values if value]
        if values and all(len(value) <= 9 and " " not in value for value in values):
            center.add(index)
    return center


def row_fills(body: Sequence[Sequence[str]], class_index: Optional[int]) -> Dict[int, str]:
    fills: Dict[int, str] = {}
    if class_index is None:
        return fills
    for index, row in enumerate(body):
        label = row[class_index].replace("**", "").strip()
        if label == "FP":
            fills[index] = AMBER
        elif label == "FN":
            fills[index] = RED
    return fills


# ---------------------------------------------------------------------------
# Document rendering
# ---------------------------------------------------------------------------


def build_cover(doc: Document, title: str, date: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    cell_margins(table, 220, 300, 220, 300)
    cell = table.cell(0, 0)
    shade(cell, NAVY)
    no_split(table.rows[0])

    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = rgb(WHITE)

    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(date)
    run.font.size = Pt(10)
    run.font.color.rgb = rgb("D9E2F3")
    fixed(table, [PORTRAIT_WIDTH])


def render_table(doc: Document, header: Sequence[str], body: Sequence[Sequence[str]], width: float) -> None:
    widths = column_widths(header, body, width)
    size = TABLE_SIZES.get(len(header), 7.6)
    wide = len(header) >= 11
    class_index = None
    for index, label in enumerate(header):
        if label.strip().lower() == "kelas":
            class_index = index
    add_table(
        doc,
        header,
        body,
        widths,
        size=size,
        header_size=6.0 if wide else size + 0.1,
        center=centered_columns(header, body),
        first_bold=len(header) <= 6,
        zebra=True,
        row_fills=row_fills(body, class_index),
        margin=(25, 45, 25, 45) if wide else None,
        keep_rows_together=len(body) <= 8,
    )


def build_document() -> Tuple[Document, Dict[str, Any]]:
    if not MD_PATH.exists():
        raise SystemExit(f"Sumber markdown tidak ditemukan: {MD_PATH}")

    blocks = parse_markdown(MD_PATH)
    if len(blocks) < 3 or blocks[0][0] != "h1":
        raise SystemExit("Markdown harus dimulai dengan judul '# ...' dan baris tanggal.")

    title = str(blocks[0][1])
    date = str(blocks[1][1])

    doc = Document()
    section = doc.sections[0]
    page_setup(section, landscape=False)
    setup_header_footer(section, PORTRAIT_WIDTH, DOC_DATE)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1

    properties = doc.core_properties
    properties.title = title
    properties.subject = "Validasi internal PreHub: praktisi, ketergunaan, model, dan tabel 60 skenario"
    properties.author = "Tim PreHub GEMASTIK 2026"
    properties.keywords = "PreHub, GEMASTIK 2026, validasi, SUS, consensus gate, SUM-SCN"
    properties.comments = "Dokumen internal, bukan pilot operasional."

    build_cover(doc, title, date)
    spacer(doc, 6.0)

    landscape = False
    table_count = 0
    band_count = 1  # cover band
    big_table = None
    big_rows = big_cols = 0

    for kind, payload in blocks[2:]:
        if kind == "h2":
            text = str(payload)
            match = re.match(r"^(\d+)\.\s*(.*)$", text)
            number = match.group(1) if match else ""
            heading = match.group(2) if match else text

            if number == "5" and not landscape:
                landscape = True
                new_section = doc.add_section(WD_SECTION.NEW_PAGE)
                page_setup(new_section, landscape=True)
                setup_header_footer(new_section, LANDSCAPE_WIDTH, DOC_DATE)

            gap_points = 0 if text.startswith("1.") else 8
            if gap_points:
                lead = doc.add_paragraph("")
                lead.paragraph_format.space_before = Pt(0)
                lead.paragraph_format.space_after = Pt(gap_points)
                lead.paragraph_format.line_spacing = Pt(1)

            section_band(
                doc,
                f"{number}." if number else "",
                heading,
                LANDSCAPE_WIDTH if landscape else PORTRAIT_WIDTH,
                size=12.0 if not landscape else 11.5,
            )
            spacer(doc, 4.0, keep_next=True)
            band_count += 1
        elif kind == "h3":
            subsection(doc, str(payload))
        elif kind == "p":
            text = str(payload)
            legend = text.startswith(("SS = ", "TP = "))
            para(
                doc,
                text,
                size=8.6 if legend else 9.0,
                italic=legend,
                color=MID if legend else None,
                space_after=3 if legend else 6,
            )
        elif kind == "table":
            header, body = payload
            render_table(doc, header, body, LANDSCAPE_WIDTH if landscape else PORTRAIT_WIDTH)
            spacer(doc, 5.0)
            table_count += 1
            if len(header) == 11:
                big_table = doc.tables[-1]
                big_rows = len(body)
                big_cols = len(header)

    stats = {
        "tables": len(doc.tables),
        "content_tables": table_count,
        "layout_bands": band_count,
        "n60_rows": big_rows,
        "n60_columns": big_cols,
        "n60_table_rows": len(big_table.rows) if big_table is not None else 0,
        "blocks": len(blocks),
        "landscape_sections": 1 if landscape else 0,
    }
    return doc, stats


# ---------------------------------------------------------------------------
# Export, verification, manifest
# ---------------------------------------------------------------------------


def backup(path: Path) -> Optional[Path]:
    if not path.exists():
        return None
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    candidate = path.with_name(f"{path.name}.bak-{stamp}")
    counter = 1
    while candidate.exists():
        candidate = path.with_name(f"{path.name}.bak-{stamp}-{counter}")
        counter += 1
    shutil.copy2(path, candidate)
    return candidate


def export_pdf_word(docx_path: Path, pdf_path: Path) -> Tuple[int, str]:
    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore

        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = None
        try:
            document = word.Documents.Open(str(docx_path), False, False, False)
            document.Fields.Update()
            document.Repaginate()
            pages = int(document.ComputeStatistics(2))
            document.SaveAs(str(pdf_path), FileFormat=17)
            document.Close(False)
            document = None
        finally:
            if document is not None:
                document.Close(False)
            word.Quit()
            pythoncom.CoUninitialize()
        return pages, "word_com"
    except Exception as word_error:
        candidates = [
            shutil.which("soffice"),
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        soffice = next((item for item in candidates if item and Path(item).exists()), None)
        if soffice:
            import subprocess

            outdir = pdf_path.parent
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=180,
            )
            produced = outdir / (docx_path.stem + ".pdf")
            if produced.exists() and produced != pdf_path:
                shutil.move(str(produced), str(pdf_path))
            return count_pdf_pages(pdf_path), "libreoffice"
        raise RuntimeError(f"Word COM gagal dan LibreOffice tidak tersedia: {word_error}")


def count_pdf_pages(pdf_path: Path) -> int:
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(pdf_path)).pages)
    except Exception:
        try:
            import fitz  # type: ignore

            return fitz.open(str(pdf_path)).page_count
        except Exception:
            return 0


def docx_text(docx_path: Path) -> str:
    from docx import Document as ReopenDocument

    document = ReopenDocument(str(docx_path))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def verify_text(docx_path: Path, pdf_path: Path) -> Dict[str, Any]:
    text = docx_text(docx_path)
    normalised = re.sub(r"\s+", " ", text)

    required = [
        DOC_TITLE,
        DOC_DATE,
        "1. Identitas",
        "2. Validasi Praktisi (5 Mitra) dan Matriks Kesesuaian Fitur",
        "2.1 Matriks kesesuaian fitur",
        "3. Uji Ketergunaan: Sepuluh Kriteria SUS dan Durasi Keputusan",
        "3.1 Sepuluh kriteria SUS dan aspek PreHub yang diuji",
        "3.2 Skor per responden",
        "3.3 Efisiensi durasi keputusan dan penyelesaian tugas",
        "4. Hasil Pengujian Akurasi Model dan Consensus Gate",
        "4.1 Ringkasan hasil pengujian deteksi dan kalibrasi model",
        "4.2 Arti metrik",
        "4.3 Bobot Empat Sinyal Konsensus dan Pemetaan Lokasi Berita",
        "4.4 Ringkasan Enam Kasus Batas (3 Alarm Palsu & 3 Kebobolan)",
        "5. Tabel Lengkap 60 Skenario Pengujian Ground Truth (SUM-SCN-001 s.d. SUM-SCN-060)",
        "PT Bumi Sawita Nusantara Berjaya",
        "21 Oktober 2026",
        "33 bernilai Sesuai atau Sangat Sesuai (82,5%)",
        "Skor SUS 85,5 (Grade A / Excellent)",
        "85,50",
        "855,0",
        "14,8 menit",
        "1,8 menit",
        "91,43% (32/35)",
        "91,67% (11/12)",
        "0,0893",
        "0,0609",
        "SUM-SCN-004",
        "SUM-SCN-048",
    ]
    missing = [token for token in required if token not in normalised]
    missing_ids = [
        f"SUM-SCN-{index:03d}"
        for index in range(1, 61)
        if f"SUM-SCN-{index:03d}" not in normalised
    ]

    forbidden = [
        "v2026-10-28-N60",
        "95fc6a1",
        "Build ",
        "Batas klaim",
        "Keterbatasan",
        "Integritas artefak",
        "SHA-256",
        "Blok verifikasi",
        "Ringkasan eksekutif",
        "Lampiran N=60",
        "Dokumen Validasi Internal & Lampiran",
    ]
    present_forbidden = [token for token in forbidden if token in normalised]

    pdf_text = ""
    try:
        from pypdf import PdfReader

        pdf_text = "\n".join((page.extract_text() or "") for page in PdfReader(str(pdf_path)).pages)
    except Exception:
        pdf_text = ""

    pdf_normalised = re.sub(r"\s+", " ", pdf_text)
    pdf_missing = [token for token in ("PreHub", "SUM-SCN-001", "SUM-SCN-060", "Lembar Validasi Internal")
                   if token not in pdf_normalised]

    return {
        "missing": missing + missing_ids,
        "forbidden": present_forbidden,
        "pdf_missing": pdf_missing,
        "pdf_ids": pdf_text.count("SUM-SCN-"),
        "docx_chars": len(text),
    }


def build_documents() -> Dict[str, Any]:
    doc, stats = build_document()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    backups = []
    for path in (DOCX_CANON, DOCX_LAMP, PDF_CANON, PDF_LAMP):
        saved = backup(path)
        if saved:
            backups.append(str(saved))

    doc.save(str(DOCX_CANON))
    shutil.copyfile(DOCX_CANON, DOCX_LAMP)

    pages, engine = export_pdf_word(DOCX_CANON, PDF_CANON)
    shutil.copyfile(PDF_CANON, PDF_LAMP)

    docx_hash = sha256_file(DOCX_CANON)
    lamp_docx_hash = sha256_file(DOCX_LAMP)
    pdf_hash = sha256_file(PDF_CANON)
    lamp_pdf_hash = sha256_file(PDF_LAMP)
    if docx_hash != lamp_docx_hash:
        raise RuntimeError("DOCX kanonik dan Lampiran tidak byte-identik.")
    if pdf_hash != lamp_pdf_hash:
        raise RuntimeError("PDF kanonik dan Lampiran tidak byte-identik.")

    verification = verify_text(DOCX_CANON, PDF_CANON)
    if verification["missing"]:
        raise RuntimeError(f"Teks wajib tidak ditemukan: {verification['missing']}")
    if verification["forbidden"]:
        raise RuntimeError(f"Teks lama yang dilarang masih muncul: {verification['forbidden']}")
    if verification["pdf_missing"]:
        raise RuntimeError(f"PDF kehilangan teks wajib: {verification['pdf_missing']}")
    if pages <= 0:
        raise RuntimeError("Jumlah halaman PDF tidak terbaca.")

    manifest = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "pdf_engine": engine,
        "source": {
            "markdown": {
                "path": str(MD_PATH.relative_to(ROOT)).replace("\\", "/"),
                "sha256": sha256_file(MD_PATH),
            }
        },
        "document": {
            "title": DOC_TITLE,
            "date": DOC_DATE,
            "tables": stats["tables"],
            "markdown_tables": stats["content_tables"],
            "layout_bands": stats["layout_bands"],
            "n60_rows": stats["n60_rows"],
            "n60_columns": stats["n60_columns"],
            "pdf_pages": pages,
            "pdf_ids_found": verification["pdf_ids"],
            "docx_chars": verification["docx_chars"],
        },
        "files": [
            {
                "path": str(DOCX_CANON.relative_to(ROOT)).replace("\\", "/"),
                "bytes": DOCX_CANON.stat().st_size,
                "sha256": docx_hash,
                "kind": "docx",
            },
            {
                "path": str(DOCX_LAMP.relative_to(ROOT)).replace("\\", "/"),
                "bytes": DOCX_LAMP.stat().st_size,
                "sha256": lamp_docx_hash,
                "kind": "docx",
            },
            {
                "path": str(PDF_CANON.relative_to(ROOT)).replace("\\", "/"),
                "bytes": PDF_CANON.stat().st_size,
                "sha256": pdf_hash,
                "kind": "pdf",
                "pages": pages,
            },
            {
                "path": str(PDF_LAMP.relative_to(ROOT)).replace("\\", "/"),
                "bytes": PDF_LAMP.stat().st_size,
                "sha256": lamp_pdf_hash,
                "kind": "pdf",
                "pages": pages,
            },
        ],
        "backups": backups,
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"Saved: {DOCX_CANON.relative_to(ROOT)} ({DOCX_CANON.stat().st_size:,} bytes)")
    print(f"Saved: {DOCX_LAMP.relative_to(ROOT)} ({DOCX_LAMP.stat().st_size:,} bytes, sha256 identical)")
    print(f"Saved: {PDF_CANON.relative_to(ROOT)} ({PDF_CANON.stat().st_size:,} bytes, {pages} pages, {engine})")
    print(f"Saved: {PDF_LAMP.relative_to(ROOT)} ({PDF_LAMP.stat().st_size:,} bytes, sha256 identical)")
    print(
        f"Tables: {stats['tables']} total ({stats['content_tables']} markdown tables + "
        f"{stats['layout_bands']} layout bands) · "
        f"N=60 table: {stats['n60_rows']} data rows x {stats['n60_columns']} columns"
    )
    print(f"PDF IDs found: {verification['pdf_ids']} · DOCX chars: {verification['docx_chars']:,}")
    print(f"Manifest: {MANIFEST_PATH.relative_to(ROOT)}")
    return manifest


if __name__ == "__main__":
    build_documents()
