"""Targeted patch: add Fase 37 marker to P206 + Table3 R14, preserve SDT."""
import hashlib, pathlib
from docx import Document
SRC = pathlib.Path("Final-Gemastik/Laporan-PreHub-GEMASTIK19.docx")
print("SHA_PATCH_BEFORE=" + hashlib.sha256(SRC.read_bytes()).hexdigest())
d = Document(str(SRC))
# P206: append Fase 37 note (SDT-safe: extend runs[0], clear others? NO — that would rebuild.
# Instead append to LAST non-empty run to keep order, SDTs at end stay after? SDTs are between runs?
# Safest: append note to runs[0] existing text? That moves note to front. Better: find last run with text and append.
para = d.paragraphs[206]
runs = [r for r in para.runs if r.text.strip() != ""]
print("P206 non-empty runs:", len(runs), "total runs:", len(para.runs))
print("P206 tail:", repr(para.text[-120:]))
note = " Konsolidasi CPU ini berlaku sejak Fase 37."
if "Fase 37" not in para.text:
    runs[-1].text = runs[-1].text + note
    print("P206 patched")
else:
    print("P206 already has Fase 37")
# Table3 R14C2
t3 = d.tables[3]
cell = t3.rows[14].cells[2]
print("R14C2 before:", repr(cell.text[:160]))
if "Fase 37" not in cell.text:
    pruns = [r for r in cell.paragraphs[0].runs if r.text.strip() != ""]
    pruns[-1].text = pruns[-1].text + " Berlaku sejak Fase 37."
    print("R14C2 patched")
else:
    print("R14C2 already has Fase 37")
d.save(str(SRC))
print("SHA_PATCH_AFTER=" + hashlib.sha256(SRC.read_bytes()).hexdigest())
print("P206 after tail:", repr(d.paragraphs[206].text[-160:]))
print("R14C2 after:", repr(d.tables[3].rows[14].cells[2].text[:220]))
