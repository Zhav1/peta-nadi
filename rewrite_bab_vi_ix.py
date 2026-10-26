"""Rewrite BAB VI-IX in Laporan-PreHub-GEMASTIK19.docx preserving Mendeley SDT, images, tables."""
import docx
from docx.shared import Pt, RGBColor
from lxml import etree

PATH = "Final-Gemastik/Laporan-PreHub-GEMASTIK19.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}

def qn(tag):
    return f"{{{W_NS}}}{tag}"

def set_plain_para(para, new_text):
    """Replace text of a paragraph without SDT/images. Preserves style and first-run font."""
    # save font attrs from first run if any
    font_name, font_size, bold, italic, color = None, None, None, None, None
    if para.runs:
        r0 = para.runs[0]
        try:
            font_name = r0.font.name
            font_size = r0.font.size
            bold = r0.font.bold
            italic = r0.font.italic
            color = r0.font.color.rgb if r0.font.color and r0.font.color.rgb else None
        except Exception:
            pass
    # Setter clears runs (safe only when no SDT/image — caller guarantees)
    para.text = ""
    run = para.add_run(new_text)
    try:
        if font_name:
            run.font.name = font_name
        if font_size:
            run.font.size = font_size
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic
        if color is not None:
            run.font.color.rgb = color
    except Exception:
        pass

def set_sdt_para(para, before_text, after_text):
    """Replace text around preserved w:sdt elements. Keeps all SDT, drops old w:r/proofErr."""
    el = para._element
    sdts = el.findall("w:sdt", NS)
    assert sdts, "set_sdt_para called on para without SDT"
    # capture formatting from first w:r
    first_r = el.find("w:r", NS)
    rpr_xml = None
    if first_r is not None:
        rpr = first_r.find("w:rPr", NS)
        if rpr is not None:
            rpr_xml = etree.tostring(rpr)
    # remove all w:r and w:proofErr and w:hyperlink children (keep pPr and sdt)
    for child in list(el):
        if child.tag in (qn("r"), qn("proofErr"), qn("hyperlink"), qn("bookmarkStart"), qn("bookmarkEnd")):
            el.remove(child)
    def make_run(text):
        r = etree.Element(qn("r"))
        if rpr_xml is not None:
            try:
                r.append(etree.fromstring(rpr_xml))
            except Exception:
                pass
        t = etree.SubElement(r, qn("t"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = text
        return r
    first_sdt = sdts[0]
    last_sdt = sdts[-1]
    idx_first = list(el).index(first_sdt)
    # insert before run
    if before_text:
        el.insert(idx_first, make_run(before_text))
        idx_last = list(el).index(last_sdt)
        if after_text:
            el.insert(idx_last + 1, make_run(after_text))
    else:
        if after_text:
            idx_last = list(el).index(last_sdt)
            el.insert(idx_last + 1, make_run(after_text))

def main():
    doc = docx.Document(PATH)
    edited = []

    # ---- Non-SDT normal paras ----
    plain = {
    169: "PreHub dirancang untuk Operations Coordinator atau Logistics Dispatcher yang memantau perjalanan distribusi dan menentukan respons ketika kondisi berubah. Pengguna membutuhkan jawaban atas lima pertanyaan operasional: apa yang terjadi, seberapa kuat buktinya, apakah perjalanan terdampak, seberapa besar dampaknya, dan apa yang dapat dilakukan. Tiga peran berlaku di aplikasi: Dispatcher memegang approvals:write, fleet:write, routes:plan; Regulator membaca heatmap, harga, laporan, audit dan memberi flag HOLD sebagai advis; Guest hanya membaca benchmark dan laporan di sandbox (backend/app/auth/supabase_auth.py:50-70). Keputusan akhir tetap di operator; persetujuan di dasbor tidak menggerakkan armada sendiri (backend/app/auth/supabase_auth.py:42-45).",
    171: "Kebutuhan tersebut diterjemahkan menjadi fungsi utama PreHub: integrasi data multisumber, deteksi indikasi gangguan, validasi bukti, estimasi probabilitas, analisis dampak, penilaian risiko, optimasi alternatif mitigasi, penyajian Evidence Chain, serta pencatatan keputusan dan hasil aktual. Fungsi berjalan sebagai pipeline deterministik 8-node di LangGraph (agents/graph.py:56-63). Sistem tidak menggantikan TMS, navigasi, atau kewenangan operator. PreHub adalah lapisan decision support yang mengubah informasi terfragmentasi menjadi konteks operasional.",
    172: "Gambar arsip Figma berikut merekam persona hasil proses emphatize. Berkas Figma bersifat arsip desain, bukan tangkapan localhost; tangkapan localhost diberi label host dan tanggal di BAB VIII.",
    175: "Gambar arsip Figma berikut merangkum use case canvas dari analisis dan persona. Sama seperti persona, canvas ini arsip desain dan tidak diklaim sebagai tangkapan live.",
    179: "PreHub menggabungkan data terstruktur (cuaca BMKG, lalu lintas TomTom, posisi armada) dan tidak terstruktur (berita, OSINT). Tahap ingestion melakukan normalisasi, validasi, deduplikasi, geokoding, dan pemetaan indeks spasial H3 dengan metadata agar dapat ditelusuri. TomTom memakai kunci berurutan key0 lalu key1; latch global 900 detik hanya aktif setelah semua kunci habis, latch berbagi Redis selalu False sehingga tiap proses mundur independen (backend/app/adapters/tomtom_adapter.py:38-58,168-171,183-189,257-288). Tanpa kunci, fetch mengembalikan degraded/tomtom_api_key_missing; saat latch aktif, fetch short-circuit (backend/app/adapters/tomtom_adapter.py:503-517). Armada menggabung baseline 45-unit di kode dengan armada kustom lewat webhook TMS yang live (backend/app/services/telemetry_service.py:239-249; backend/app/routers/fleet_ingest_router.py:266-286). Berita dan OSINT diubah Intelligence Agent menjadi kejadian terstruktur berisi lokasi, waktu, jenis, dan bukti.",
    182: "PreHub membedakan indikasi gangguan dari kejadian yang telah tervalidasi. Tiap agen domain (cuaca, lalu lintas, OSINT, ekonomi) mengeluarkan sinyal berisi sumber, waktu, lokasi, jenis bukti, kualitas, dan keyakinan. Consensus Engine mengumpulkan sinyal dan menghitung skor gabungan, bukan menghitung jumlah sumber. Aturan dedup: status duplicate langsung ke END agar payload sama tidak diingest ulang; sumber operator, UI, dan simulasi dikecualikan agar demo dan replay tidak tertelan TTL 300 detik (agents/graph.py:32-41).",
    183: "Validasi tidak menghitung sumber secara naif karena BMKG dan model atmosfer berbagi fenomena yang sama. PreHub mengelompokkan bukti per domain dan menilai independensi, kualitas, kebaruan, kesesuaian spasial, dan konsistensi. Temuan optimasi rute internal (Agent-4) tidak ikut voting sesuai FR-11.2 (agents/tools/consensus_gate.py:9-10,405-406). Kanal ekonomi (PIHPS) hanya voting saat anomali aktif econ_conf > 0,50; bila tidak, bobot fisik dinormalisasi ulang W 0,4375 / T 0,3125 / I 0,25 / E 0,0 agar krisis fisik tidak terdiskon 20 persen (agents/tools/consensus_gate.py:28-44,368-371). Bukti atmosfer, lalu lintas, dan laporan lapangan yang konsisten memiliki nilai validasi lebih kuat dibanding beberapa keluaran model dari fenomena yang sama.",
    185: "PreHub memisahkan evidence confidence dari probabilitas gangguan. Keduanya berbeda dan diuji dengan metrik berbeda.",
    186: "Evidence confidence menilai dukungan bukti terhadap indikasi: kualitas sumber, kelengkapan, kebaruan, kesesuaian spasial-temporal, dan konsistensi. Confidence berjangkar cuaca (level peringatan, presipitasi, risiko banjir, angin, kode BMKG, severity payload), bukan kesehatan sumber di Supabase yang hanya diagnostik (agents/nodes/data_collection.py:246).",
    187: "Probabilitas gangguan menilai peluang kondisi tersebut benar-benar mengganggu koridor yang diamati. Nilai tidak diambil dari confidence model bahasa atau jumlah sumber. Alur konseptual: Evidence lalu Validation lalu Disruption Probability. Rumus independensi probabilistik yang dipakai: P_disruption(s) = 1 - prod_k (1 - w_k(t,d) * p_k(s)) untuk k di {W,T,I,E} (agents/tools/consensus_gate.py:3-4). Bobot dasar saat anomali aktif: W 0,35 / T 0,25 / I 0,20 / E 0,20; saat tidak aktif: W 0,4375 / T 0,3125 / I 0,25 / E 0,0 (agents/tools/consensus_gate.py:29-44). Tiap bobot dikali peluruhan temporal e^(-0,05 * jam) dan spasial e^(-d / 25 km) (agents/tools/consensus_gate.py:49-50,81-88). Suku per kanal: term_k = 1 - w_k * p_k; peluang non-deteksi = term_W * term_T * term_I * term_E; P mentah = 1 - peluang non-deteksi (agents/tools/consensus_gate.py:396-403).",
    188: "Probabilitas diuji terhadap outcome aktual. Kalibrasi memakai P_max = 1 - prod(1 - w_k) agar konfirmasi multisensor penuh mencapai 1,0; fungsi calibrate_confidence_scale membagi P mentah dengan P_max (agents/tools/consensus_gate.py:91-108). Contoh: kelompok prediksi 0,70 yang kejadian aktualnya sekitar 70 persen menunjukkan kalibrasi baik. Skala ganda berlaku: NFR-4 menilai skala ternormalisasi, produk mentah disimpan sebagai raw_brier_score; deteksi tetap di 0,50 mentah (setara 0,727 ternormalisasi); overall_confidence tidak diskorkan karena cap OSINT 0,50 akan memotong benchmark (.planning/LEARNINGS.md:276).",
    189: "Evaluasi probabilitas memakai Brier Score dan analisis kalibrasi; precision, recall, F1 menilai klasifikasi. Nilai probabilitas dibandingkan dengan outcome, bukan diterima karena keluar dari model. Laporan kanonik sinkron 2026-10-26: N=60, TP=35, FP=0, TN=25, FN=0, Brier=0,0937 (backend/app/routers/evaluation_router.py:32-56).",
    191: "PreHub tidak menentukan risiko dari confidence saja. Kejadian probabilitas tinggi dengan dampak rendah tidak menjadi prioritas; kejadian probabilitas sedang dengan keterlambatan besar di koridor kritis dapat menjadi prioritas. Penilaian menggabung probabilitas dan dampak: estimasi keterlambatan, kepentingan koridor, eksposur armada, severity, dan ketersediaan alternatif.",
    194: "Bentuk fungsi, bobot, dan ambang kategori ditentukan dari data dan diuji lewat skenario, bukan nilai universal. Threshold dibandingkan dengan outcome untuk mengidentifikasi false alarm, missed disruption, dan salah prioritas. Dampak moneter dihitung lewat matriks hedging: susut = CargoValue * (1 - e^(-delta * t)) (backend/app/services/spoilage_hedging_service.py:316); tier delta 0,025 per jam ultra-perishable (cabai, tomat, sayur), semi-perishable untuk umbi dan bawang, dry bulk untuk beras, minyak, gula (backend/app/services/spoilage_hedging_service.py:133-160); net_savings = cost_continue - biaya kebijakan optimal (backend/app/services/spoilage_hedging_service.py:547); solve_hedging_matrix memilih CONTINUE, REROUTE, atau HOLD (backend/app/services/spoilage_hedging_service.py:357).",
    195: "Jika probabilitas tinggi tetapi dampak rendah, sistem tidak memberi prioritas intervensi tinggi. Sebaliknya, kejadian probabilitas lebih rendah tetap prioritas bila konsekuensi besar dan alternatif terbatas. Perbandingan Continue-Reroute-Hold selalu memakai estimasi waktu, kelayakan constraint, dan eksposur risiko yang sama.",
    197: "PreHub membagi analitik menjadi komponen spesialis per domain dan jenis tugas. Implementasi adalah pipeline deterministik 8-node: data_collection, osint_hazard, prediction, consensus_gate, economic_intelligence, route_optimization, decision_support, archive (agents/graph.py:56-63). Empat tahap berlaku: ingestion, observasi sensor paralel ke fan-in gate, asesmen dampak kondisional, lalu optimasi dan keputusan sekuensial (agents/graph.py:44-52). Istilah swarm pada dokumen lama berarti orkestrasi pipeline spesialis ini.",
    199: "Pembagian dipakai karena tiap domain memiliki sumber, metode, dan keluaran berbeda. Agen tidak memutuskan final secara independen; tiap agen mengeluarkan structured output untuk tahap berikut. Duplikat berhenti di Stage 1; sumber UI dan simulasi dikecualikan agar demo tidak tertelan dedup (agents/graph.py:32-41).",
    203: "Dalam PreHub, LangGraph berfungsi sebagai orchestrator, bukan sumber kebenaran. Validasi evidence, perhitungan metrik, constraint routing, dan keputusan akhir tetap di lapisan eksplisit yang dapat diperiksa. Uji kompilasi graph dan gate 0,85 dengan 2 sensor mengunci perilaku ini (backend/app/routers/evaluation_router.py:120-122).",
    205: "Setelah gangguan tervalidasi dan relevan dengan perjalanan, PreHub memetakan lokasi, koridor, perubahan ETA, kendaraan terekspos, dan alternatif. Tepi graph yang memotong hazard diberi penalti 50x inti, 15x zona dalam, 3,5x buffer luar (backend/app/adapters/cpu_routing_adapter.py:719-726). Koridor intermodal memakai rantai truk-pelabuhan-kapal dengan perimeter laut 29 waypoint agar kapal tidak memotong darat (backend/app/adapters/cpu_routing_adapter.py:230-289,315-361).",
    207: "Keluaran optimasi tidak diterapkan ke armada. PreHub membandingkan alternatif dengan indikator waktu, kelayakan constraint, dan eksposur risiko. Tiga alternatif utama ke operator adalah:",
    211: "Operator memilih berdasar informasi sistem dan konteks lapangan. Solver CPU adalah optimization engine; PreHub adalah lapisan keputusan yang menghubungkan hasil optimasi dengan bukti dan konteks. Tidak ada eksekusi otomatis; tombol mitigasi membuka persetujuan HITL yang dicatat di SQLite (backend/app/db/local_storage.py:370; backend/app/routers/fleet_ingest_router.py:266-286).",
    215: "Evaluasi bertingkat tiga: komponen, integrasi, end-to-end. Tiap modul diukur sesuai karakter: akurasi dan latensi deteksi, ketepatan ekstraksi OSINT, Brier untuk probabilistik, efisiensi komputasi routing. Matriks uji 133 kasus FR-1 sampai FR-20 menjadi acuan (backend/app/routers/evaluation_router.py:81-217).",
    216: "Integrasi menguji aliran data hingga rekomendasi dan stabilitas saat data terlambat atau bertentangan. End-to-end memakai skenario logistik dan mengukur dampak ke keputusan. Hasil menjadi dasar iterasi parameter dan ambang. Satu catatan merah terbuka: benchmark TP 35 versus ekspektasi lama 34; label dan cut diperbaiki dulu, ekspektasi tidak diubah diam-diam (.planning/LEARNINGS.md:290).",
    218: "PreHub adalah decision-support closed-loop dari pantau hingga evaluasi. Data multisumber dinormalisasi menjadi bukti; modul cuaca, lalu lintas, OSINT menganalisis; Consensus Engine memvalidasi berdasar konsistensi, kebaruan, dan independensi. Sistem mengukur probabilitas dan dampak tanpa bertumpu pada satu keyakinan. Gate promosi: confidence >= 0,85 dan >=2 sensor independen plus koroborasi fisik atau bypass simulasi (agents/tools/consensus_gate.py:482-489). Klaim tak terkoroborasi dibatasi <=0,50; bukti simulasi didiskon 0,90, bukan ditulis ulang menjadi kepastian (agents/tools/consensus_gate.py:473-480).",
    221: "Jika risiko butuh tindakan, modul optimasi menyusun Continue, Reroute, Hold beserta bukti hedging dan proof rute. Pilihan tampil transparan lewat Evidence Chain. Keputusan di operator; tindakan dan hasil dicatat untuk evaluasi. Rekam dipakai menyempurnakan parameter dan analitik sehingga siklus deteksi hingga perbaikan berjalan berlanjut. Persistensi lokal memakai SQLite prehub_local.db agar jejak tetap ada saat Supabase atau Redis terganggu (backend/app/db/local_storage.py:17-19,34).",
    224: "PreHub adalah aplikasi web modular: Next.js di frontend, FastAPI di backend, PostgreSQL dengan PostGIS untuk spasial. Nama produk PreHub; alias deploy peta-nadi (render.yaml:3; render.yaml:69-72). URL publik dan CORS memakai https://peta-nadi.onrender.com; SOP lokal memakai http://localhost:3000 dan http://localhost:8000. Struktur modular memudahkan uji komponen terpisah dan penggantian fixture simulasi dengan sumber operasional bertahap tanpa mengganggu antarmuka. Status operasional: demonstrator decision-support untuk koridor pangan Sumatra, bukan sistem operasional kritis nasional; belum ada pilot bersama BULOG, Bapanas, atau operator armada (Final-Gemastik/last-judgement.txt:11). Label BULOG di kode adalah fixture (backend/app/services/telemetry_service.py:247; backend/app/routers/fleet_ingest_router.py:49), bukan bukti integrasi.",
    228: "Ingestion menormalisasi, memvalidasi, mendedup, geocode, dan memetakan H3. Event diteruskan lewat event stream dengan metadata. Probe live 26-Okt memakai scripts/check_tomtom.py:40-89: mask kunci, cek Redis, status latch, matriks klasifier 6 kasus, health dan fetch live kondisional. Perilaku live: kunci berurutan dengan failover key0 ke key1, latch 900 detik setelah semua kunci habis; tanpa kunci, status degraded (scripts/check_tomtom.py:61-65; backend/app/adapters/tomtom_adapter.py:503-507). Armada: baseline 45-unit di kode difusi dengan armada kustom; endpoint webhook TMS live (backend/app/routers/fleet_ingest_router.py:266-286). Split hybrid 99/73 adalah observasi probe 26-Okt terhadap /api/v1/fleet/vehicles untuk diverifikasi ulang, bukan konstanta kode; laporan tidak meng-hardcode angka tersebut ke tabel. Data yang belum real-time diganti simulasi berlabel agar tidak campur dengan operasional.",
    230: "Lapisan analitik memakai komponen spesialis cuaca, lalu lintas, publik, risiko, dan alternatif. Graph 8-node deterministik mengatur status dan konteks tanpa menjadikan keluaran bahasa sebagai kebenaran tunggal (agents/graph.py:56-91). Agent-4 (optimasi rute) tidak voting di gate (agents/tools/consensus_gate.py:9-10). Kanal E hanya voting saat econ_conf > 0,50 (agents/tools/consensus_gate.py:368-371). Bukti simulasi didiskon 0,90 (agents/tools/consensus_gate.py:52-55). Modul yang dirancang dapat memakai simulasi tanpa mengganggu alur utama.",
    232: "Consensus Engine menilai sinyal berdasar kualitas, kebaruan, spasial-temporal, dan independensi sebelum menjadi keputusan. Status peringatan memakai ambang terkalibrasi dari data operasional dan metrik. Syarat valid: >=0,85, >=2 sensor, dan koroborasi fisik (telemetri BMKG terkonfirmasi atau delay TomTom >=0,50 dalam <=30 km) atau bypass simulasi (agents/tools/consensus_gate.py:464-485). Validasi silang BMKG deterministik memakai toleransi spasial <=30 km (<=50 km gempa M>=4,5) dan temporal <=6 jam dengan ambang fisik 20 mm per jam, SIAGA atau WASPADA, M>=4,5, gelombang >=2,5 m (agents/tools/consensus_gate.py:147-254).",
    233: "Jika data kurang, bertentangan, atau sumber utama terganggu, sistem degrade dengan provenance dan menurunkan keyakinan atau meminta verifikasi operator. Kegagalan satu sumber tidak menghentikan sistem dan tidak memicu rekomendasi tanpa bukti kuat. Jarak hilang diperlakukan sebagai tak diketahui (inf) agar tidak lolos gate <=30 km (agents/tools/consensus_gate.py:257-274).",
    235: "PreHub memvalidasi info perjalanan bersama sumber lain sebelum menyimpulkan gangguan. Prakiraan cuaca diperlakukan sebagai sinyal masa depan yang dipetakan ke koridor; akurasi cuaca diuji terpisah dari deteksi. Fallback koridor memakai radius 25 km dan delay per severity saat cache TomTom kosong (agents/nodes/prediction.py:24-29,288-316).",
    236: "Mitigasi menyusun Continue, Reroute, Hold berdasar kendala. Alternatif tidak dieksekusi otomatis; disajikan dengan waktu, risiko, dan kelayakan. Solve hedging memilih kebijakan optimal dan net_savings; proof rute memuat dasar risiko, uncertainty dampak, dan keterbatasan data (backend/app/adapters/cpu_routing_adapter.py:59-223,874-909).",
    238: "Dasbor menyatukan koridor, bukti, risiko, dan mitigasi dalam satu alur. Evidence Chain menelusuri sumber hingga rekomendasi. Mitigasi menampilkan Continue, Reroute, Hold/Delay dengan waktu dan risiko. Dispatcher menyetujui; Regulator hanya memberi flag HOLD di drawer HITL tanpa approvals:write; Guest hanya membuka pratinjau sandbox ber-watermark tanpa POST (backend/app/auth/supabase_auth.py:50-70). Penolakan server 403 berlaku untuk Regulator dan Guest pada /approvals dan fleet:write pada /fleet/register, /fleet/upload-manifest, /fleet/custom (backend/app/auth/supabase_auth.py:386-393; backend/tests/test_auth_rbac.py:233-251; backend/app/routers/fleet_ingest_router.py:108-120). Tidak ada eksekusi otomatis.",
    240: "Status komponen dibedakan agar rancangan tidak dibaca sebagai sudah terimplementasi semua. Produk adalah demonstrator decision-support, bukan early-warning operasional; infrastruktur gratis cocok untuk demo, tidak untuk SLA, redundansi, keamanan, dan kapasitas kritis (Final-Gemastik/last-judgement.txt:11). Tidak ada klaim kesiapan kritis nasional dan tidak ada aktuasi otomatis (Final-Gemastik/last-judgement.txt:11,20). Tabel VII.1 dipertahankan apa adanya; penjelasan status live versus fallback ada di paragraf ini agar tabel tidak diubah.",
    243: "Pengujian bertahap dari fungsi dasar, API, aliran data, hingga skenario utama. Modul simulasi diuji konsistensi alur tanpa klaim performa nyata. Metrik per modul: akurasi dan latensi deteksi, Brier, ketepatan ekstraksi, kelayakan rute dan waktu komputasi. Laporan kanonik adalah fallback statis sinkron 2026-10-26 (N=60, TP=35, FP=0, TN=25, FN=0, Brier 0,0937) yang dipakai hanya saat test-results/benchmark_evaluation_report.json tidak terbaca (backend/app/routers/evaluation_router.py:30-56). Endpoint mock dan eval berlabel fallback; angka tidak diklaim live bila JSON kanonik tidak terbaca.",
    244: "Skenario end-to-end menguji ketahanan saat data tidak lengkap atau bertentangan. Pengujian pengguna (SUS, task completion, decision time, error rate, baseline manual) ditunda; UI padat dan belum diuji dengan dispatcher nyata (Final-Gemastik/last-judgement.txt:13). Kuantifikasi kerugian koridor untuk urgensi juga menunggu (Final-Gemastik/last-judgement.txt:22). Hasil menjadi dasar iterasi parameter, ambang, dan antarmuka.",
    256: "Penggunaan dimulai saat Koordinator Operasional memantau koridor di dasbor localhost:3000 (Playwright baseURL dan webServer localhost:3000; frontend/playwright.config.ts:14,25-32). Galeri tangkapan docs/screenshots 01 sampai 07 terikat ke Gambar 7.1 sampai 7.7 (scripts/generate_docx_technical_doc.py:541-597). Saat gangguan terdeteksi, sistem memberi peringatan berisi lokasi, jenis, risiko. Dispatcher memverifikasi lewat Evidence Chain, membuka Impact Analysis untuk koridor dan ETA. Regulator membaca dan memberi flag HOLD di drawer HITL; Guest membuka pratinjau sandbox ber-watermark. Semua tangkapan berlabel host dan tanggal; tidak dilabeli produksi.",
    259: "Selanjutnya pengguna membuka Mitigation Recommendation untuk membandingkan Continue, Reroute, Hold/Delay dengan waktu, risiko, dan konsekuensi. Dispatcher memilih dan POST /approvals; Regulator dan Guest ditolak 403 di server (backend/app/auth/supabase_auth.py:386-393; backend/tests/test_auth_rbac.py:233-251). HOLD adalah flag advis HITL, bukan commit approvals:write (backend/app/auth/supabase_auth.py:42-45). PreHub tidak mengubah rute otomatis. LIVE_GPS hanya saat ping GPS atau TMS, AIS LIVE, atau ADS-B LIVE ada (backend/app/services/telemetry_service.py:1294,1521,1634); selain itu ROUTE_ESTIMATED dengan SIMULATION_CACHE dan sumber ROUTE_DEAD_RECKONING (backend/app/services/telemetry_service.py:1652-1654,1771-1773). Cakupan upstream: AISstreamAdapter untuk maritim, airplanes.live dan OpenSky untuk udara, webhook TMS untuk truk; tanpa upstream, baseline 45-unit tampil sebagai fallback berlabel (backend/app/services/telemetry_service.py:1889-1916). Setelah keputusan, pantau kondisi dan bandingkan aktual versus prediksi untuk evaluasi.",
    261: "PreHub membedakan indikasi awal dan kejadian tervalidasi lewat status dan risiko. Sumber dan confidence ditelusuri di Evidence Chain. Data tidak lengkap atau sumber terganggu ditampilkan dan tidak dihitung sebagai bukti positif; jarak hilang bernilai inf (agents/tools/consensus_gate.py:257-274). Berita fallback berlabel dan dikecualikan dari indikator live. Operator dapat menerima, menolak, atau override; keputusan dicatat di SQLite prehub_local.db (backend/app/db/local_storage.py:370) untuk audit dan evaluasi. Tanpa token, 401 mendahului 422 (backend/tests/test_auth_rbac.py:233-251); tanpa approvals:write, server mengembalikan 403 (backend/app/auth/supabase_auth.py:386-393).",
    }
    for idx, txt in plain.items():
        try:
            set_plain_para(doc.paragraphs[idx], txt)
            edited.append(idx)
        except Exception as e:
            print(f"PLAIN P{idx} FAIL: {e}")

    # ---- Heading-5 captions (no SDT) ----
    captions = {
    174: "User Persona (Arsip Figma, bukan tangkapan localhost)",
    177: "Use Case Canvas (Arsip Figma, bukan tangkapan localhost)",
    202: "Agentic Workflow Diagram (Arsip desain 8-node deterministik; istilah swarm lama berarti orkestrasi pipeline)",
    220: "End-to-end Flow Diagram (Arsip desain closed-loop; gate 0,85 dengan 2 sensor)",
    226: "Diagram Arsitektur (Arsip desain; produk PreHub, alias deploy peta-nadi)",
    247: "Tampilan Hero Section PreHub (localhost:3000, 26-Okt-2026; LIVE bila API dan Redis up, label SIMULASI atau OFFLINE bila fallback; bukan produksi)",
    249: "Tampilan Evidence Chain (localhost:3000, 26-Okt-2026; LIVE bila bukti live, fallback berlabel dan dikecualikan dari indikator live)",
    251: "Tampilan Antarmuka Map 4D PreHub (localhost:3000, 26-Okt-2026; LIVE_GPS hanya saat ping, AIS, atau ADS-B ada, selain itu ROUTE_ESTIMATED dengan SIMULATION_CACHE)",
    253: "Tampilan Antarmuka Simulasi Kebijakan Alternatif Rute (localhost:3000, 26-Okt-2026; Continue-Reroute-Hold dengan hedging; HOLD advis, bukan commit)",
    258: "User Flow Diagram (localhost:3000; alur Dispatcher-Regulator-Guest; HOLD UI-only, server 403)",
    }
    for idx, txt in captions.items():
        try:
            set_plain_para(doc.paragraphs[idx], txt)
            edited.append(idx)
        except Exception as e:
            print(f"CAPTION P{idx} FAIL: {e}")

    # ---- SDT paras: preserve citations ----
    sdt_edits = {
    180: ("Untuk prakiraan cuaca, desain PreHub merujuk FourCastNet berbasis Adaptive Fourier Neural Operator ",
            " sebagai referensi model. Implementasi yang berjalan memakai fusi Open-Meteo (ECMWF dan GFS) dengan BMKG tanpa GPU; routing terkonsolidasi di CPU NetworkX dan OR-Tools dengan cache jalan lokal (backend/app/adapters/cpu_routing_adapter.py:1-5; .planning/PROJECT.md:51). Seluruh data terintegrasi diolah menjadi objek bukti tiga lapis: data mentah, bukti terstruktur, dan konteks operasional, sehingga kualitas sumber dan hasil analisis dinilai independen dan data belum lengkap tidak langsung menentukan keputusan."),
    200: ("Orkestrasi memakai LangGraph sebagai runtime workflow stateful berisi node, edge, branching, persistence, dan human-in-the-loop ",
            ". LangGraph adalah orchestrator, bukan sumber kebenaran; validasi evidence, metrik, constraint routing, dan keputusan akhir berada di lapisan eksplisit yang dapat diperiksa (agents/graph.py:68-91). Istilah swarm pada dokumen lama berarti orkestrasi pipeline spesialis, bukan kolaborasi otonom (Final-Gemastik/last-judgement.txt:8)."),
    206: ("Untuk alternatif rute, desain PreHub merujuk NVIDIA cuOpt sebagai mesin optimasi ",
            ". Implementasi yang berjalan adalah Deterministic CPU Solver: NetworkX Dijkstra dan A* untuk detour point-to-point (<30 ms) dan VRP CPU NetworkX dengan heuristik 2-Opt (<150 ms per 50 node), nol GPU (backend/app/adapters/cpu_routing_adapter.py:1-5,607-627,927-936; .planning/PROJECT.md:51). Label solver: NetworkX Dijkstra / A* CPU Solver dan Deterministic CPU VRP Solver (NetworkX + Heuristic 2-Opt) (backend/app/adapters/cpu_routing_adapter.py:911-913,990-992). Cabang fallback rute eksplisit tersedia bila tidak ada path (backend/app/adapters/cpu_routing_adapter.py:818-870). Permasalahan rute direpresentasikan lewat lingkungan perjalanan, armada, tugas, waktu, dan kendala seperti kapasitas, time windows, prioritas, dan batas operasi."),
    213: ("PreHub menyediakan Evidence Chain dari sumber, analisis, validasi, hingga rekomendasi. Rekam relevan: lokasi, waktu, keyakinan, pemicu. Saat konflik atau data terbatas, sistem degrade dengan provenance: PIHPS fallback 14 hari commodity_prices, STM dan dedup lokal terbatas 500 entri, OSINT 12 jam berita terakhir, prediksi SCAN dengan anchor koridor 25 km (agents/nodes/data_collection.py:369; agents/nodes/prediction.py:24-29). Sinyal arsip memaksa validated=false, membuka Evidence, dan toast info agar UI tidak memalsu sukses; berita fallback berlabel dan dikecualikan dari indikator live. Operator dapat mengoreksi rekomendasi; riwayat tindakan dicatat untuk audit. Desain mendukung transparansi dan interpretabilitas sesuai panduan NIST AI RMF ",
            ""),
    }
    for idx, (bef, aft) in sdt_edits.items():
        try:
            set_sdt_para(doc.paragraphs[idx], bef, aft)
            edited.append(idx)
        except Exception as e:
            print(f"SDT P{idx} FAIL: {e}")

    doc.save(PATH)
    print(f"EDITED_PARAS: {sorted(edited)}")
    print(f"COUNT: {len(edited)}")

if __name__ == "__main__":
    main()
